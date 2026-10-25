"""Shared grading engine for the CH 227z skill checks.

Each skill check is a small config dict (see spectroscopy.py) that names its
inputs, its per-version ranges and constants, and an ordered list of steps
built from the reusable step types below. This module reads the TA's input
sheet (.csv or .xlsx), runs the steps for every student, raises flags for
manual review and writes a graded copy in the same column layout.

Each skill check lives in its own "N Name Skill Check" folder next to this
file and shares roster.csv. Adding one means writing a new config in that
folder and calling run(config, folder).
"""

from __future__ import annotations

import argparse
import csv
import html
import math
import sys
import webbrowser
from datetime import datetime
from pathlib import Path

YES, NO = "Yes", "No"
HERE = Path(__file__).resolve().parent
ROSTER = HERE / "roster.csv"

NAME, NUMBER, ABSENT = "Student Name", "Number", "Absent"
VERSION, CHEATING, FLAGS = "Version (A/B)", "Cheating?", "Flags"
ABSENT_MARKS = {"y", "yes", "x", "true", "1", "absent"}


# --- Version rules -----------------------------------------------------------

def parity_version(number):
    """Odd student number -> Version A, even -> Version B."""
    return "A" if number % 2 else "B"


# --- Step types --------------------------------------------------------------
# Every step has a key (so later steps can use its result), an output column
# header, and evaluate(values, config, version). A step returns None when an
# input it depends on is missing, so the cell stays blank instead of "No".

class RangeCheck:
    """Yes when low <= value <= high for the student's version (inclusive)."""

    def __init__(self, key, column, field):
        self.key, self.column, self.field = key, column, field

    def bounds(self, config, version):
        return config["ranges"][version][self.field]

    def evaluate(self, values, config, version):
        value = values.get(self.field)
        if value is None:
            return None
        low, high = self.bounds(config, version)
        return YES if low <= value <= high else NO


class Computed:
    """A derived number, e.g. the expected answer from a calibration curve."""

    def __init__(self, key, column, func, requires):
        self.key, self.column, self.func, self.requires = key, column, func, requires

    def evaluate(self, values, config, version):
        if any(values.get(field) is None for field in self.requires):
            return None
        return self.func(values, config)


class ToleranceCheck:
    """Yes when |student - expected| <= tolerance * |expected| (relative)."""

    def __init__(self, key, column, field, expected):
        self.key, self.column, self.field, self.expected = key, column, field, expected

    def evaluate(self, values, config, version):
        student, expected = values.get(self.field), values.get(self.expected)
        if student is None or expected is None:
            return None
        return YES if abs(student - expected) <= abs(expected) * config["tolerance"] else NO


class AllPass:
    """Yes only if every listed check is Yes; blank until all have a result."""

    def __init__(self, key, column, checks):
        self.key, self.column, self.checks = key, column, checks

    def evaluate(self, values, config, version):
        results = [values.get(key) for key in self.checks]
        if any(result is None for result in results):
            return None
        return YES if all(result == YES for result in results) else NO


# --- Grading -----------------------------------------------------------------

def run_steps(config, inputs, version):
    """Run every step in order; returns inputs plus each step's result."""
    values = dict(inputs)
    for step in config["steps"]:
        values[step.key] = step.evaluate(values, config, version)
    return values


def review_flags(config, values, number, version):
    """Notes for the TA: inputs that fit another version, near misses, 10^n slips."""
    flags = []
    others = [v for v in config["ranges"] if v != version]
    for step in config["steps"]:
        if values.get(step.key) != NO:
            continue
        if isinstance(step, RangeCheck):
            value = values[step.field]
            label = config["inputs"][step.field]
            for other in others:
                low, high = step.bounds(config, other)
                if low <= value <= high:
                    flags.append(f"{label} {value:.3g} fits Version {other} "
                                 f"(number {number} is Version {version})")
        elif isinstance(step, ToleranceCheck):
            student, expected = values[step.field], values[step.expected]
            label = config["inputs"][step.field]
            if student == 0 or expected == 0:
                continue
            if student * expected < 0:
                flags.append(f"{label} has the wrong sign")
                continue
            ratio = student / expected
            off = abs(ratio - 1)
            power = round(math.log10(ratio))
            if off <= 2 * config["tolerance"]:
                flags.append(f"{label} is {off:.1%} off (limit "
                             f"{config['tolerance']:.0%}), near miss")
            elif power and abs(ratio / 10 ** power - 1) <= config["tolerance"]:
                flags.append(f"{label} is off by 10^{power} (unit or exponent slip?)")
    return flags


def grade_student(config, row):
    """Grade one parsed row; returns version, step results, cheating flag, notes."""
    graded = {"version": None, "values": {}, "cheating": None, "flags": list(row["flags"])}
    inputs = row["inputs"]
    entered = [v for v in inputs.values() if v is not None]
    if row["absent"]:
        graded["flags"].append("Absent, not graded")
        return graded
    if not entered:
        graded["flags"].append("No inputs entered, not graded")
        return graded
    if len(entered) == len(inputs) and all(v == 0 for v in entered):
        graded["flags"].append("All inputs are zero, treated as absent")
        return graded
    if row["number"] is None:
        graded["flags"].append("No student number, not graded")
        return graded

    number = row["number"]
    version = config["version"](number)
    values = run_steps(config, inputs, version)
    graded.update(version=version, values=values)
    missing = [config["inputs"][key] for key, value in inputs.items() if value is None]
    if missing:
        graded["flags"].append(f"Incomplete, missing {', '.join(missing)}")
    zeros = [config["inputs"][key] for key, value in inputs.items() if value == 0]
    if zeros:
        graded["flags"].append(f"{', '.join(zeros)} entered as 0 (not turned in?)")
    graded["flags"] += review_flags(config, values, number, version)

    # Cheating?: the same answers would pass if graded as another version.
    matches = [other for other in config["ranges"] if other != version
               and run_steps(config, inputs, other)[config["pass"]] == YES]
    if matches:
        graded["cheating"] = YES
        graded["flags"].append(f"Answers pass as Version {', '.join(matches)}")
    return graded


# --- Input -------------------------------------------------------------------

def parse_number(raw, label, flags):
    """Blank -> None; numbers and numeric text -> float; anything else is flagged."""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    text = raw.strip().replace("−", "-").replace(",", "")
    try:
        return float(text)
    except ValueError:
        flags.append(f"{label} {raw!r} is not a number")
        return None


def read_table(path, sheet=None):
    """Rows of {header: value} from a .csv or the named sheet of an .xlsx."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl

        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = workbook[sheet] if sheet in workbook.sheetnames else workbook.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
        table = [dict(zip(header, row)) for row in rows]
        workbook.close()
        return table
    with open(path, newline="", encoding="utf-8-sig") as f:
        return [{(k or "").strip(): v for k, v in row.items()} for row in csv.DictReader(f)]


def load_students(config, path):
    """Parse the input sheet into rows of name, number, absent flag and inputs."""
    students = []
    for raw in read_table(path, config.get("sheet")):
        name = raw.get(NAME)
        name = str(name).strip() if name is not None else ""
        if not name and raw.get(NUMBER) in (None, ""):
            continue
        flags = []
        number = parse_number(raw.get(NUMBER), NUMBER, flags)
        absent = str(raw.get(ABSENT) or "").strip().lower() in ABSENT_MARKS
        inputs = {key: parse_number(raw.get(header), header, flags)
                  for key, header in config["inputs"].items()}
        students.append({"name": name, "number": int(number) if number is not None else None,
                         "absent": absent, "inputs": inputs, "flags": flags})
    return students


# --- Output ------------------------------------------------------------------

def output_columns(config):
    """(key, header) pairs after Name/Number/Version, matching the workbook layout."""
    headers = dict(config["inputs"])
    headers.update({step.key: step.column for step in config["steps"]})
    return [(key, headers[key]) for key in config["columns"]]


def graded_rows(config, students, results):
    header = [NAME, NUMBER, VERSION] + [h for _, h in output_columns(config)] + [CHEATING, FLAGS]
    rows = []
    for student, graded in zip(students, results):
        row = [student["name"], student["number"], graded["version"]]
        for key, _ in output_columns(config):
            row.append(graded["values"].get(key, student["inputs"].get(key)))
        row += [graded["cheating"], "; ".join(graded["flags"])]
        rows.append(row)
    return header, rows


def write_csv(path, header, rows):
    def cell(value):
        if value is None:
            return ""
        if isinstance(value, float):
            return f"{value:.4g}"
        return value

    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows([cell(v) for v in row] for row in rows)


def write_xlsx(path, header, rows, sheet):
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = sheet or "Grading"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, float) and 0 < abs(cell.value) < 1e-2:
                cell.number_format = "0.00E+00"
            elif cell.value == YES:
                cell.fill = red if header[cell.column - 1] == CHEATING else green
            elif cell.value == NO:
                cell.fill = red
    ws.freeze_panes = "B2"
    workbook.save(path)


def print_summary(config, students, results):
    print(config["name"])
    print("-" * len(config["name"]))
    width = max(len(s["name"]) for s in students)
    for student, graded in zip(students, results):
        result = graded["values"].get(config["pass"]) or "-"
        mark = " CHEATING?" if graded["cheating"] else ""
        notes = "; ".join(graded["flags"])
        print(f"{student['name']:<{width}}  {graded['version'] or '-'}  "
              f"{result:<4}{mark}  {notes}".rstrip())
    passes = [g["values"].get(config["pass"]) for g in results]
    print(f"\nPassed {passes.count(YES)}, failed {passes.count(NO)}, "
          f"not graded {passes.count(None)}, cheating flags "
          f"{sum(1 for g in results if g['cheating'])}")


# --- HTML report -------------------------------------------------------------

REPORT_CSS = """
:root { color-scheme: light dark;
  --bg: #f6f7f9; --card: #ffffff; --text: #1d2330; --muted: #5d6675; --line: #dfe3ea;
  --yes: #1f7a43; --yes-bg: #ddf3e4; --no: #b4232c; --no-bg: #fbe1e3;
  --warn: #8a5300; --warn-bg: #fdf0d5; --none-bg: #eceef2; }
@media (prefers-color-scheme: dark) { :root {
  --bg: #15181e; --card: #1e222a; --text: #e6e9ef; --muted: #9aa3b2; --line: #313743;
  --yes: #7ad69b; --yes-bg: #1d3a28; --no: #ff8f96; --no-bg: #44222a;
  --warn: #f3c46b; --warn-bg: #3e3119; --none-bg: #2a2f39; } }
* { box-sizing: border-box; }
body { margin: 0; background: var(--bg); color: var(--text);
  font: 15px/1.45 "Segoe UI", system-ui, sans-serif; }
main { max-width: 1200px; margin: 0 auto; padding: 24px 16px 48px; }
h1 { font-size: 24px; margin: 0; }
h2 { font-size: 17px; margin: 32px 0 12px; }
.meta { color: var(--muted); margin: 4px 0 20px; font-size: 13px; }
.tiles { display: grid; grid-template-columns: repeat(auto-fit, minmax(150px, 1fr)); gap: 12px; }
.tile { background: var(--card); border: 1px solid var(--line); border-radius: 10px; padding: 14px 16px; }
.tile b { display: block; font-size: 30px; line-height: 1.1; font-variant-numeric: tabular-nums; }
.tile span { color: var(--muted); font-size: 13px; }
.tile.yes b { color: var(--yes); } .tile.no b { color: var(--no); } .tile.warn b { color: var(--warn); }
.review { display: grid; grid-template-columns: repeat(auto-fill, minmax(320px, 1fr)); gap: 12px; }
.card { background: var(--card); border: 1px solid var(--line); border-left: 4px solid var(--muted);
  border-radius: 10px; padding: 12px 14px; }
.card.cheat { border-left-color: var(--no); } .card.flag { border-left-color: var(--warn); }
.card header { display: flex; gap: 8px; align-items: baseline; flex-wrap: wrap; }
.card header strong { font-size: 15px; } .card header small { color: var(--muted); }
.card ul { margin: 8px 0 0; padding-left: 18px; font-size: 13.5px; }
.pill { display: inline-block; padding: 1px 8px; border-radius: 999px; font-size: 12px; font-weight: 600; }
.pill.yes { color: var(--yes); background: var(--yes-bg); }
.pill.no { color: var(--no); background: var(--no-bg); }
.pill.cheat { color: var(--no); background: var(--no-bg); }
.pill.none { color: var(--muted); background: var(--none-bg); }
.filters { display: flex; gap: 6px; flex-wrap: wrap; margin-bottom: 10px; }
.filters button { font: inherit; font-size: 13px; padding: 4px 12px; border-radius: 999px; cursor: pointer;
  border: 1px solid var(--line); background: var(--card); color: var(--text); }
.filters button.on { background: var(--text); color: var(--card); border-color: var(--text); }
.table-wrap { overflow-x: auto; background: var(--card); border: 1px solid var(--line); border-radius: 10px; }
table { border-collapse: collapse; width: 100%; font-size: 13.5px; }
th, td { padding: 7px 10px; border-bottom: 1px solid var(--line); text-align: left; vertical-align: top; }
th { position: sticky; top: 0; background: var(--card); font-size: 12px; color: var(--muted); white-space: nowrap; }
td.num { font-variant-numeric: tabular-nums; white-space: nowrap; }
td.flags { color: var(--muted); font-size: 12.5px; min-width: 420px; }
tr:last-child td { border-bottom: 0; }
tr.hidden { display: none; }
.empty { color: var(--muted); }
"""

REPORT_JS = """
document.querySelectorAll('.filters button').forEach(function (button) {
  button.addEventListener('click', function () {
    document.querySelectorAll('.filters button').forEach(function (b) { b.classList.remove('on'); });
    button.classList.add('on');
    var want = button.dataset.filter;
    document.querySelectorAll('tbody tr').forEach(function (row) {
      var show = want === 'all' || row.dataset.status === want || (want === 'flagged' && row.dataset.flagged);
      row.classList.toggle('hidden', !show);
    });
  });
});
"""


def report_number(value):
    """Short HTML number: 0.656, or 2.95 × 10<sup>-5</sup> for very small/large values."""
    if value == 0 or 1e-3 <= abs(value) < 1e4:
        return f"{value:.4g}"
    mantissa, exponent = f"{value:.2e}".split("e")
    return f"{mantissa} &times; 10<sup>{int(exponent)}</sup>"


def report_cell(value, header):
    if value is None:
        return "<td></td>"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f'<td class="num">{report_number(value)}</td>'
    if value in (YES, NO):
        kind = "cheat" if header == CHEATING else value.lower()
        return f'<td><span class="pill {kind}">{value}</span></td>'
    return f"<td>{html.escape(str(value))}</td>"


def write_report(config, students, results, path, source):
    """Self-contained HTML summary: totals, students to review, full results table."""
    passes = [g["values"].get(config["pass"]) for g in results]
    cheats = sum(1 for g in results if g["cheating"])
    tiles = [("yes", passes.count(YES), "Passed"), ("no", passes.count(NO), "Failed"),
             ("", passes.count(None), "Not graded"), ("warn", cheats, "Cheating? flags")]

    # Students to look at by hand: cheating flags first, then other notes.
    review = sorted((pair for pair in zip(students, results) if pair[1]["flags"]),
                    key=lambda pair: not pair[1]["cheating"])
    cards = []
    for student, graded in review:
        result = graded["values"].get(config["pass"])
        label = {YES: "Pass", NO: "Fail"}.get(result, "Not graded")
        badge = f'<span class="pill {result.lower() if result else "none"}">{label}</span>'
        if graded["cheating"]:
            badge += ' <span class="pill cheat">Cheating?</span>'
        version = f"Version {graded['version']}" if graded["version"] else "No version"
        notes = "".join(f"<li>{html.escape(flag)}</li>" for flag in graded["flags"])
        kind = "cheat" if graded["cheating"] else "flag" if graded["version"] else ""
        cards.append(f'<div class="card {kind}"><header><strong>{html.escape(student["name"])}</strong>'
                     f"<small>#{student['number'] if student['number'] is not None else '?'} "
                     f"&middot; {version}</small>{badge}</header><ul>{notes}</ul></div>")

    header, rows = graded_rows(config, students, results)
    head = "".join(f"<th>{html.escape(h)}</th>" for h in header)
    body = []
    for row, graded in zip(rows, results):
        result = graded["values"].get(config["pass"])
        status = {YES: "pass", NO: "fail"}.get(result, "ungraded")
        flagged = ' data-flagged="1"' if graded["flags"] else ""
        cells = [report_cell(v, h) for v, h in zip(row[:-1], header)]
        cells.append(f'<td class="flags">{html.escape(row[-1])}</td>')
        body.append(f'<tr data-status="{status}"{flagged}>{"".join(cells)}</tr>')

    filters = [("all", "All"), ("fail", "Failed"), ("pass", "Passed"),
               ("flagged", "Needs review"), ("ungraded", "Not graded")]
    page = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(config["name"])} Results</title>
<style>{REPORT_CSS}</style></head>
<body><main>
<h1>{html.escape(config["name"])}</h1>
<p class="meta">Graded {datetime.now():%B %d, %Y at %I:%M %p} from {html.escape(source.name)}
&middot; {len(students)} students &middot; tolerance &plusmn;{config.get("tolerance", 0):.0%}</p>
<div class="tiles">{"".join(f'<div class="tile {k}"><b>{n}</b><span>{label}</span></div>'
                             for k, n, label in tiles)}</div>
<h2>Needs review ({len(cards)})</h2>
<div class="review">{"".join(cards) or '<p class="empty">Nothing flagged.</p>'}</div>
<h2>All results</h2>
<div class="filters">{"".join(f'<button data-filter="{k}"{" class=on" if k == "all" else ""}>{label}</button>'
                               for k, label in filters)}</div>
<div class="table-wrap"><table><thead><tr>{head}</tr></thead>
<tbody>{"".join(body)}</tbody></table></div>
</main><script>{REPORT_JS}</script></body></html>
"""
    path.write_text(page, encoding="utf-8")


# --- Template & CLI ----------------------------------------------------------

def write_template(config, path):
    """Blank input sheet from the shared roster, in gradebook order."""
    with open(ROSTER, newline="", encoding="utf-8-sig") as f:
        roster = list(csv.DictReader(f))
    header = [NAME, NUMBER, ABSENT] + list(config["inputs"].values())
    rows = [[r[NAME], r[NUMBER]] + [""] * (len(header) - 2) for r in roster]
    write_csv(path, header, rows)


def run(config, folder, argv=None):
    """Command line entry point; folder is the skill check's own directory."""
    default_input = Path(folder) / f"{config['slug']}_input.csv"
    parser = argparse.ArgumentParser(description=f"Grade the {config['name']}.")
    parser.add_argument("input", nargs="?", type=Path, default=default_input,
                        help=f"input .csv or .xlsx (default: {default_input.name})")
    parser.add_argument("-o", "--output", type=Path,
                        help="graded copy to write (default: <input>_graded, same format)")
    parser.add_argument("--template", action="store_true",
                        help="write a blank input sheet from roster.csv and exit")
    parser.add_argument("--no-open", action="store_true",
                        help="write the HTML report without opening it in the browser")
    args = parser.parse_args(argv)
    # The Windows console defaults to cp1252, which can't print every name.
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    if args.template:
        if args.input.exists():
            parser.error(f"{args.input} already exists; not overwriting it")
        write_template(config, args.input)
        print(f"Wrote {args.input}")
        return

    students = load_students(config, args.input)
    results = [grade_student(config, student) for student in students]
    print_summary(config, students, results)

    output = args.output or args.input.with_name(f"{args.input.stem}_graded{args.input.suffix}")
    header, rows = graded_rows(config, students, results)
    if output.suffix.lower() in (".xlsx", ".xlsm"):
        write_xlsx(output, header, rows, config.get("sheet"))
    else:
        write_csv(output, header, rows)
    print(f"\nWrote {output}")

    report = output.with_suffix(".html")
    write_report(config, students, results, report, args.input)
    print(f"Wrote {report}")
    if not args.no_open:
        webbrowser.open(report.resolve().as_uri())
